"""Importing admissions from Excel: the template, reading and validating the upload, the routes."""

from __future__ import annotations

import json
import re
from datetime import datetime, timedelta
from io import BytesIO
from pathlib import Path
from typing import Any
from zoneinfo import ZoneInfo

import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook
from sqlalchemy import func, select
from sqlalchemy.orm import Session

from robby.clinical import importer
from robby.clinical.form import DetailSpec, parse_time
from robby.clinical.importer import RowProblem, find_duplicates, read_upload
from robby.clinical.scoring import ModelRuntime
from robby.clinical.xlsx_template import build_template, columns
from robby.config import Settings, default_rule_set_path
from robby.models import Admission, AdmissionSource, Classification, ModelState, Prediction, Role
from robby.rulesets import load_rule_set_file
from tests.factories import events_of, make_facility, make_model_version, make_user
from tests.test_clinical_flow import GOOD, StubScorer, now_text
from tests.webutil import csrf_for, sign_in_client

SPECS = [
    DetailSpec("height_cm", "Height", "number", "cm", 120, 200, None),
    DetailSpec("hiv_positive", "HIV positive", "category", "", None, None, ("No", "Yes")),
]
ZONE = ZoneInfo("Africa/Kigali")
NOW = datetime(2026, 10, 18, 12, 0, tzinfo=ZONE)

HEADER = [
    "Patient number",
    "Patient name",
    "Admitted at",
    "Maternal age",
    "Gestational age (weeks)",
    "Parity",
    "Previous cesareans",
    "Number of fetuses",
    "Fetal presentation",
    "Onset of labour",
    "Systolic BP",
    "Diastolic BP",
    "BP not measured: reason",
    "Proteinuria",
    "Glucose (mmol/L)",
    "HIV positive",
]


def good_row(number: str = "RW-1", admitted: object = "2026-10-18 09:30") -> list[object]:
    return [
        number, "Uwase Marie", admitted, 29, 39, 2, 1, "One", "Cephalic", "Spontaneous",
        128, 84, None, "Negative", 5.2, "yes",
    ]  # fmt: skip


def make_xlsx(header: list[object], rows: list[list[object]], title: str = "Admissions") -> bytes:
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = title
    ws.append(header)
    for row in rows:
        ws.append(row)
    out = BytesIO()
    wb.save(out)
    return out.getvalue()


def read(data: bytes, filename: str = "admissions.xlsx") -> importer.ReadResult:
    return read_upload(data, filename, SPECS, ZONE, NOW)


def problems_of(data: bytes, filename: str = "admissions.xlsx") -> list[RowProblem]:
    result = read(data, filename)
    assert result.rows == []
    return result.problems


# ---------------------------------------------------------------- template


def test_template_has_headers_lists_validation_and_help() -> None:
    wb = load_workbook(BytesIO(build_template(SPECS)))
    assert wb.sheetnames == ["Admissions", "Lists", "How to fill this in"]
    sheet = wb["Admissions"]
    header = [c.value for c in next(sheet.iter_rows(min_row=1, max_row=1))]
    assert header == [c.label for c in columns(SPECS)]
    assert header[:3] == ["Patient number", "Patient name", "Admitted at"]
    assert header[-2:] == ["Height", "HIV positive"]
    lists = wb["Lists"]
    listed = {
        col[0].value: [c.value for c in col[1:] if c.value is not None]
        for col in lists.iter_cols()
        if col[0].value
    }
    assert listed["Fetal presentation"] == ["Cephalic", "Breech", "Transverse", "Oblique"]
    assert listed["HIV positive"] == ["No", "Yes"]
    assert "Height" not in listed
    presentation = header.index("Fetal presentation") + 1
    letter = sheet.cell(row=1, column=presentation).column_letter
    ranges = [str(dv.sqref) for dv in sheet.data_validations.dataValidation]
    assert any(r.startswith(f"{letter}2") for r in ranges)
    assert wb["How to fill this in"]["A1"].value


def test_column_fields_are_the_form_field_names() -> None:
    fields = [c.field for c in columns(SPECS)]
    assert fields[:3] == ["patient_number", "patient_name", "admitted_at"]
    assert {"bp_sys", "bp_dia", "bp_reason", "proteinuria", "protein_reason"} <= set(fields)
    assert {"glucose", "glucose_reason", "d_height_cm", "d_hiv_positive"} <= set(fields)
    assert [c.label for c in columns(SPECS) if c.required] == [
        "Patient number",
        "Patient name",
        "Admitted at",
    ]


def test_the_template_itself_reads_back_with_no_rows() -> None:
    assert problems_of(build_template(SPECS)) == [
        RowProblem(0, "", "The file has no admissions to import.")
    ]


# ---------------------------------------------------------------- reading


def test_a_valid_file_gives_parsed_rows() -> None:
    data = make_xlsx(HEADER, [good_row("RW-1"), good_row("RW-2", datetime(2026, 10, 17, 22, 5))])
    result = read(data)
    assert result.problems == []
    assert [n for n, _ in result.rows] == [2, 3]
    first, second = (p for _, p in result.rows)
    assert first.patient_number == "RW-1" and first.patient_name == "Uwase Marie"
    assert first.admitted_at == parse_time("2026-10-18 09:30", ZONE)
    assert second.admitted_at == datetime(2026, 10, 17, 22, 5, tzinfo=ZONE)
    assert first.presentation == "cephalic" and first.plurality == 1
    assert first.maternal_age == 29 and first.ga_weeks == 39
    assert first.details == {"hiv_positive": "Yes"}  # normalised to the trained spelling


def test_the_sheet_named_admissions_is_preferred() -> None:
    wb = Workbook()
    first = wb.active
    assert first is not None
    first.title = "Notes"
    first.append(["not", "a", "header"])
    sheet = wb.create_sheet("Admissions")
    sheet.append(HEADER)
    sheet.append(good_row())
    out = BytesIO()
    wb.save(out)
    assert len(read(out.getvalue()).rows) == 1


def test_file_level_problems() -> None:
    good = make_xlsx(HEADER, [good_row()])
    assert problems_of(good, "admissions.csv") == [RowProblem(0, "", "Upload an Excel .xlsx file.")]
    assert problems_of(b"not a workbook") == [
        RowProblem(0, "", "The file could not be read as an Excel workbook.")
    ]
    many = make_xlsx(HEADER, [good_row(f"RW-{i}") for i in range(importer.MAX_ROWS + 1)])
    assert problems_of(many) == [RowProblem(0, "", "At most 500 rows can be imported at once.")]


def test_header_problems() -> None:
    unknown = make_xlsx([*HEADER[:5], "Parity!", *HEADER[6:]], [good_row()])
    assert problems_of(unknown) == [
        RowProblem(1, "Parity!", "Unknown column “Parity!”. Use the headers from the template.")
    ]
    no_name = make_xlsx([h for h in HEADER if h != "Patient name"], [])
    assert problems_of(no_name) == [
        RowProblem(1, "Patient name", "The column “Patient name” is missing.")
    ]


def test_headers_match_ignoring_case_and_spacing() -> None:
    header = [h.upper().replace(" ", "  ") for h in HEADER]
    assert len(read(make_xlsx(header, [good_row()])).rows) == 1


def test_blank_rows_are_skipped() -> None:
    data = make_xlsx(HEADER, [good_row("RW-1"), [None] * len(HEADER), good_row("RW-2")])
    assert [n for n, _ in read(data).rows] == [2, 4]


def test_row_problems_name_the_row_and_column() -> None:
    bad = good_row()
    bad[0] = None
    bad[10] = None  # systolic blank, diastolic given: the BP error
    data = make_xlsx(HEADER, [good_row(), bad])
    assert problems_of(data) == [
        RowProblem(3, "Patient number", "Enter the patient number."),
        RowProblem(
            3,
            "Systolic BP",
            "Blood pressure needs both numbers, as whole numbers (for example 128 / 84).",
        ),
    ]


def test_a_reason_stands_in_for_a_missing_value() -> None:
    row = good_row()
    row[10], row[11], row[12] = None, None, "woman declined"
    result = read(make_xlsx(HEADER, [row]))
    assert result.problems == []
    parsed = result.rows[0][1]
    assert parsed.bp_reason == "declined" and parsed.systolic is None


def test_a_value_and_a_reason_together_is_a_problem() -> None:
    row = good_row()
    row[12] = "declined"
    assert problems_of(make_xlsx(HEADER, [row])) == [
        RowProblem(
            2, "BP not measured: reason", "Enter a value or a reason it was not measured, not both."
        )
    ]


def test_cell_text() -> None:
    assert importer.cell_text(None) == ""
    assert importer.cell_text(32.0) == "32"
    assert importer.cell_text(5.2) == "5.2"
    assert importer.cell_text(7) == "7"
    assert importer.cell_text("  RW-1 ") == "RW-1"
    assert importer.cell_text(datetime(2026, 10, 18, 9, 30, 15)) == "2026-10-18 09:30"


# ---------------------------------------------------------------- duplicates


def _existing(db: Session, facility: Any, number: str, at: datetime, **kw: Any) -> Admission:
    user = make_user(db, facility=facility)
    admission = Admission(
        facility_id=facility.id,
        code=f"{facility.code}-{number}-{kw.get('code', '1')}",
        admitted_at=at,
        source=AdmissionSource.FORM,
        entered_by=user.id,
        patient_number=number,
        patient_name="Test Patient",
        supersedes_id=kw.get("supersedes"),
        correction_reason="fix" if kw.get("supersedes") else None,
    )
    db.add(admission)
    db.flush()
    return admission


def test_duplicates_in_the_file_and_in_the_facility(db: Session) -> None:
    facility = make_facility(db)
    other = make_facility(db)
    at = datetime(2026, 10, 18, 9, 30, tzinfo=ZONE)
    stored = _existing(db, facility, "RW-7", at)
    old = _existing(db, facility, "RW-8", at)
    # corrected to another admission time: only the replaced row matches, and it does not count
    _existing(db, facility, "RW-8", at.replace(hour=10), supersedes=old.id, code="2")
    _existing(db, other, "RW-9", at)
    rows = read(
        make_xlsx(
            HEADER,
            [
                good_row("RW-1"),
                good_row("RW-1"),
                good_row("RW-7"),
                good_row("RW-9"),
                good_row("RW-8"),
            ],
        )
    ).rows
    assert len(rows) == 5
    scan = find_duplicates(db, facility.id, rows)
    assert scan.in_file == [
        RowProblem(3, "Patient number", "Same patient and admission time as row 2.")
    ]
    assert {n: a.code for n, a in scan.recorded.items()} == {4: stored.code}


# ---------------------------------------------------------------- routes

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


@pytest.fixture
def clinic(app: FastAPI, client: TestClient, db: Session, settings: Settings) -> dict[str, Any]:
    load_rule_set_file(db, default_rule_set_path())
    facility = make_facility(db)
    user = make_user(db, role=Role.CLINICIAN, facility=facility)
    row = sign_in_client(client, db, settings, user)
    make_model_version(db, state=ModelState.ACTIVE)
    db.commit()
    app.state.model_runtime = ModelRuntime(
        Path("unused"), loader=lambda path, rule_set=None: StubScorer()
    )
    recent = (datetime.now(settings.zone) - timedelta(hours=2)).strftime("%Y-%m-%d %H:%M")
    return {"csrf": csrf_for(settings, row), "facility": facility, "admitted": recent}


def upload(client: TestClient, clinic: dict[str, Any], data: bytes, **form: str) -> Any:
    return client.post(
        "/admissions/import",
        data={"csrf_token": clinic["csrf"], **form},
        files={"file": ("admissions.xlsx", data, XLSX)},
    )


def admission_count(db: Session) -> int:
    return db.scalar(select(func.count()).select_from(Admission)) or 0


def test_import_page_and_template_download(client: TestClient, clinic: dict[str, Any]) -> None:
    page = client.get("/admissions/import")
    assert page.status_code == 200
    assert 'href="/admissions/import/template.xlsx"' in page.text
    assert 'type="file" name="file"' in page.text and 'enctype="multipart/form-data"' in page.text
    assert '<a class="on" aria-current="page" href="/admissions/import">' in page.text
    template = client.get("/admissions/import/template.xlsx")
    assert template.status_code == 200
    assert template.headers["content-type"] == XLSX
    assert template.headers["content-disposition"].startswith("attachment;")
    header = next(load_workbook(BytesIO(template.content))["Admissions"].iter_rows(max_row=1))
    assert [c.value for c in header][-2:] == ["Height", "HIV positive"]


def test_a_valid_upload_saves_every_row(
    client: TestClient, clinic: dict[str, Any], db: Session
) -> None:
    rows = [good_row("RW-1", clinic["admitted"]), good_row("RW-2", clinic["admitted"])]
    response = upload(client, clinic, make_xlsx(HEADER, rows))
    assert response.status_code == 200
    assert "2 admissions imported" in response.text
    saved = list(db.scalars(select(Admission).order_by(Admission.code)))
    assert [(a.patient_number, a.source) for a in saved] == [
        ("RW-1", AdmissionSource.IMPORT),
        ("RW-2", AdmissionSource.IMPORT),
    ]
    assert all(f'href="/admissions/{a.code}"' in response.text for a in saved)
    assert len(list(db.scalars(select(Classification)))) == 2
    assert len(list(db.scalars(select(Prediction)))) == 2
    imported = events_of(db, "admissions.imported")
    assert len(imported) == 1
    payload = imported[0].payload
    assert payload["facility"] == clinic["facility"].code and payload["rows"] == 2
    assert len(payload["file_sha256"]) == 64
    created = events_of(db, "admission.created")
    assert [e.payload["source"] for e in created] == ["import", "import"]
    assert all("RW-" not in json.dumps(e.payload) for e in [*created, *imported])


def test_one_bad_row_saves_nothing(client: TestClient, clinic: dict[str, Any], db: Session) -> None:
    bad = good_row("", clinic["admitted"])
    response = upload(
        client, clinic, make_xlsx(HEADER, [good_row("RW-1", clinic["admitted"]), bad])
    )
    assert response.status_code == 400
    assert "Nothing was imported" in response.text
    assert "Enter the patient number." in response.text and "<td>3</td>" in response.text
    assert admission_count(db) == 0
    assert events_of(db, "admissions.imported") == []


def test_a_repeated_upload_shows_the_existing_records_and_adds_nothing(
    client: TestClient, clinic: dict[str, Any], db: Session
) -> None:
    data = make_xlsx(HEADER, [good_row("RW-1", clinic["admitted"])])
    first = upload(client, clinic, data)
    assert first.status_code == 200
    code = db.scalar(select(Admission.code))
    again = upload(client, clinic, data)
    assert again.status_code == 200
    assert "0 admissions imported, 1 already recorded" in again.text
    assert f'href="/admissions/{code}"' in again.text and "already recorded</span>" in again.text
    assert "gnum" in again.text  # the Robson group is shown, as for a new row
    assert admission_count(db) == 1
    assert len(events_of(db, "admissions.imported")) == 1  # the second upload saved nothing


def test_a_mixed_upload_imports_the_new_rows_and_lists_both_in_file_order(
    client: TestClient, clinic: dict[str, Any], db: Session
) -> None:
    upload(client, clinic, make_xlsx(HEADER, [good_row("RW-1", clinic["admitted"])]))
    rows = [good_row("RW-2", clinic["admitted"]), good_row("RW-1", clinic["admitted"])]
    response = upload(client, clinic, make_xlsx(HEADER, rows))
    assert (
        response.status_code == 200 and "1 admission imported, 1 already recorded" in response.text
    )
    codes = re.findall(r'href="/admissions/(DA-\d+|[A-Z]+-\d+)"', response.text)
    assert len(codes) >= 2 and admission_count(db) == 2


def test_a_row_repeated_inside_one_file_is_still_refused(
    client: TestClient, clinic: dict[str, Any], db: Session
) -> None:
    row = good_row("RW-1", clinic["admitted"])
    response = upload(client, clinic, make_xlsx(HEADER, [row, row]))
    assert (
        response.status_code == 400 and "Same patient and admission time as row 2" in response.text
    )
    assert admission_count(db) == 0


def test_upload_needs_a_file_and_a_csrf_token(
    client: TestClient, clinic: dict[str, Any], db: Session
) -> None:
    no_file = client.post("/admissions/import", data={"csrf_token": clinic["csrf"]})
    assert no_file.status_code == 400 and "Choose a file to upload." in no_file.text
    data = make_xlsx(HEADER, [good_row("RW-1", clinic["admitted"])])
    forged = client.post("/admissions/import", files={"file": ("admissions.xlsx", data, XLSX)})
    assert forged.status_code == 403
    assert admission_count(db) == 0


def test_save_admission_marks_only_imports_in_the_event(
    client: TestClient, clinic: dict[str, Any], db: Session, settings: Settings
) -> None:
    form = {**GOOD, "admitted_at": now_text(settings)}
    assert (
        client.post("/admissions", data={**form, "csrf_token": clinic["csrf"]}).status_code == 303
    )
    assert "source" not in events_of(db, "admission.created")[0].payload
