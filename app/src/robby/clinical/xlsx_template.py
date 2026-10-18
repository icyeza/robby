"""The Excel template clinicians download to import admissions in bulk (one row per admission).
Columns and their allowed values come from the same labels the admission form uses."""

from __future__ import annotations

from dataclasses import dataclass
from io import BytesIO

from openpyxl import Workbook
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from robby.clinical import labels
from robby.clinical.form import DetailSpec

SHEET = "Admissions"
TEMPLATE_ROWS = 500  # rows the drop-down lists cover (the import limit)


@dataclass(frozen=True)
class Column:
    label: str  # header text in row 1
    field: str  # form field name the value is passed to parse() under
    required: bool = False
    options: tuple[tuple[str, str], ...] | None = None  # (stored value, label shown)


def columns(specs: list[DetailSpec]) -> list[Column]:
    """Fixed admission columns, then one column per optional model input of the active model."""
    fixed = [
        Column("Patient number", "patient_number", required=True),
        Column("Patient name", "patient_name", required=True),
        Column("Admitted at", "admitted_at", required=True),
        Column("Maternal age", "maternal_age"),
        Column("Gestational age (weeks)", "gestational_age_weeks"),
        Column("Parity", "parity"),
        Column("Previous cesareans", "previous_cs_count"),
        Column("Number of fetuses", "plurality", options=labels.PLURALITY_OPTIONS),
        Column("Fetal presentation", "fetal_presentation", options=labels.PRESENTATION_OPTIONS),
        Column("Onset of labour", "onset_of_labour", options=labels.ONSET_OPTIONS),
        Column("Systolic BP", "bp_sys"),
        Column("Diastolic BP", "bp_dia"),
        Column("BP not measured: reason", "bp_reason", options=labels.REASONS["bp"]),
        Column("Proteinuria", "proteinuria", options=labels.PROTEINURIA_OPTIONS),
        Column(
            "Proteinuria not measured: reason", "protein_reason", options=labels.REASONS["protein"]
        ),
        Column("Glucose (mmol/L)", "glucose"),
        Column("Glucose not measured: reason", "glucose_reason", options=labels.REASONS["glucose"]),
    ]
    detail = [
        Column(s.label, f"d_{s.name}", options=tuple(s.options) if s.levels else None)
        for s in specs
    ]
    return fixed + detail


HELP_LINES = (
    "One row per admission, starting on row 2 of the Admissions sheet. Do not rename the headers.",
    "Patient number, Patient name and Admitted at are required on every row.",
    "Admitted at: a date and time within the last 30 days, e.g. 2026-10-18 14:20.",
    "Blood pressure, proteinuria and glucose: enter a value, or leave it blank and choose a "
    "reason in the matching 'not measured' column.",
    "Choice columns: pick from the drop-down list.",
    "If any row has a problem, nothing is imported; the page lists every problem to fix.",
    f"At most {TEMPLATE_ROWS} rows per file.",
)


def build_template(specs: list[DetailSpec]) -> bytes:
    cols = columns(specs)
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = SHEET
    ws.append([c.label for c in cols])
    lists = wb.create_sheet("Lists")
    for i, col in enumerate(cols, start=1):
        letter = get_column_letter(i)
        ws.column_dimensions[letter].width = max(14, len(col.label) + 2)
        if not col.options:
            continue
        lists.cell(row=1, column=i, value=col.label)
        for r, (_, shown) in enumerate(col.options, start=2):
            lists.cell(row=r, column=i, value=shown)
        last = len(col.options) + 1
        dv = DataValidation(
            type="list", formula1=f"=Lists!${letter}$2:${letter}${last}", allow_blank=True
        )
        ws.add_data_validation(dv)
        dv.add(f"{letter}2:{letter}{TEMPLATE_ROWS + 1}")
    ws.freeze_panes = "A2"
    help_ws = wb.create_sheet("How to fill this in")
    for line in HELP_LINES:
        help_ws.append([line])
    out = BytesIO()
    wb.save(out)
    return out.getvalue()
